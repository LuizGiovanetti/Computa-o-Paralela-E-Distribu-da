import os
import pandas as pd
import matplotlib.pyplot as plt
from openpyxl.styles import Font, Alignment

TXT_PARALELO = "resultados.txt"
TXT_SERIAL   = "resultados_serial.txt"
XLSX         = "resultados.xlsx"
PNG          = "grafico.png"

if not os.path.exists(TXT_SERIAL):
    raise SystemExit(
        f"Arquivo '{TXT_SERIAL}' nao encontrado. "
        "Rode o kmeans_serial.exe primeiro."
    )

df_serial = pd.read_csv(TXT_SERIAL, sep="\t", comment="#")
tempo_serial = float(df_serial.loc[0, "tempo_serial"])

print(f"Tempo serial de referencia: {tempo_serial:.6f} s")

if not os.path.exists(TXT_PARALELO):
    raise SystemExit(
        f"Arquivo '{TXT_PARALELO}' nao encontrado. "
        "Rode o kmeans_paralelo.exe primeiro."
    )

df = pd.read_csv(TXT_PARALELO, sep="\t", comment="#")
df = df.sort_values("Nucleos").reset_index(drop=True)

df["Speed-Up"]       = tempo_serial / df["Tempo de Execucao (s)"]
df["Speed-Up Ideal"] = df["Nucleos"]
df["Eficiencia"]     = df["Speed-Up"] / df["Nucleos"]

tabela = pd.DataFrame({
    "Núcleos":                df["Nucleos"].astype(int),
    "Tempo de Execução (s)":  df["Tempo de Execucao (s)"].round(3),
    "Speed-Up":               df["Speed-Up"].round(1),
    "Speed-Up Ideal":         df["Speed-Up Ideal"].astype(int),
    "Eficiência":             df["Eficiencia"].round(1),
})

print("\n===== TABELA DE RESULTADOS =====")
print(tabela.to_string(index=False))

with pd.ExcelWriter(XLSX, engine="openpyxl") as writer:
    tabela.to_excel(writer, sheet_name="Resultados", index=False)

    ws = writer.sheets["Resultados"]

    larguras = {"A": 10, "B": 22, "C": 12, "D": 16, "E": 12}
    for col, largura in larguras.items():
        ws.column_dimensions[col].width = largura

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in ws.iter_rows(min_row=2, min_col=1, max_col=1):
        for cell in row:
            cell.alignment = Alignment(horizontal="center")

    for row in ws.iter_rows(min_row=2, min_col=2, max_col=5):
        for cell in row:
            cell.number_format = "0.0"
            cell.alignment = Alignment(horizontal="center")

print(f"\nPlanilha '{XLSX}' gerada.")

fig, ax1 = plt.subplots(figsize=(10, 6))

x = tabela["Núcleos"].values
largura = 0.6

ax1.bar(x, tabela["Eficiência"], width=largura, color="#E8A33D",
        label="Eficiência", zorder=2)
ax1.set_xlabel("Núcleos", fontsize=12)
ax1.set_ylabel("Eficiência", fontsize=12)
ax1.set_ylim(0, 1.05)
ax1.set_xticks(x)

ax2 = ax1.twinx()
ax2.plot(x, tabela["Speed-Up"], "o-", color="#1F4E79",
         linewidth=2, markersize=8, label="Speed-up", zorder=3)
ax2.plot(x, tabela["Speed-Up Ideal"], "s--", color="#2E8B57",
         linewidth=1.5, markersize=6, label="Speed-up Ideal", zorder=3)
ax2.set_ylabel("Razão de Aceleração", fontsize=12)

limite = max(tabela["Speed-Up Ideal"].max(), tabela["Speed-Up"].max()) * 1.05
ax2.set_ylim(0, limite)

ax1.set_title("Aplicação / Máquina", fontsize=14, fontweight="bold")

linhas1, labels1 = ax1.get_legend_handles_labels()
linhas2, labels2 = ax2.get_legend_handles_labels()
ax1.legend(linhas1 + linhas2, labels1 + labels2,
           loc="lower center", ncol=3, bbox_to_anchor=(0.5, -0.18),
           frameon=False)

plt.tight_layout()
plt.savefig(PNG, dpi=150, bbox_inches="tight")
print(f"Grafico '{PNG}' gerado.")

plt.show()